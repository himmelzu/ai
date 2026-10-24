#!/usr/bin/env python3
"""Generate editable proposal files from the reviewed Korean Markdown draft.

Requires python-docx==1.2.0. Run from any working directory:
    python scripts/create_proposal_docs.py

The DOCX uses ordinary paragraphs and a table, without text boxes or images.
Author-only submission notes are deliberately omitted from both exports.
"""
from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "공모전_제안서_초안.md"
OUTPUT = ROOT / "downloads"
FONT = "맑은 고딕"


def set_font(run, *, size=None, bold=None):
    run.font.name = FONT
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + name), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def inline(paragraph, text, *, size=None, all_bold=False):
    for index, part in enumerate(re.split(r"\*\*(.*?)\*\*", text)):
        if part:
            run = paragraph.add_run(part)
            set_font(run, size=size, bold=all_bold or index % 2 == 1)


def text_only(text):
    return text.replace("**", "").replace("`", "")


def set_cell_margins(cell):
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", "85"), ("bottom", "85"), ("left", "100"), ("right", "100")):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:w"), value)
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    properties.append(margins)


def add_table(document, rows):
    table = document.add_table(rows=0, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Cm(2.0), Cm(7.8), Cm(7.8))
    for column, width in zip(table.columns, widths):
        column.width = width
    for row_index, values in enumerate(rows):
        row = table.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for cell, width, value in zip(row.cells, widths, values):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.12
            paragraph.paragraph_format.space_after = Pt(0)
            if row_index == 0 or cell is row.cells[0]:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(paragraph, value, size=9.2, all_bold=row_index == 0)
            if row_index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EDF5F2")
                cell._tc.get_or_add_tcPr().append(shade)
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
    return table


def build():
    content = SOURCE.read_text(encoding="utf-8").split("\n---\n", 1)[0].strip()
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.7)
    section.header_distance = section.footer_distance = Cm(0.7)
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.2)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.line_spacing = 1.26
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name in ("Heading 1", "Heading 2", "Title"):
        document.styles[name].font.name = FONT
        document.styles[name]._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    props = document.core_properties
    proposal_title = next((text_only(line).removeprefix("제안명:").strip()
                           for line in content.splitlines() if text_only(line).startswith("제안명:")),
                          "아이디어 제안서")
    props.title = proposal_title + " — 아이디어 제안서 초안"
    props.subject = "직원 AI 활용 업무 아이디어 공모전"
    props.author = ""
    props.last_modified_by = ""
    props.comments = "편집 가능한 초안입니다. 개인 참가정보와 제출 전 사실관계는 작성자가 확인하세요."
    lines = content.splitlines()
    plain = []
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            plain.append("")
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = lines[i].strip()
                i += 1
                if re.fullmatch(r"[|\s:-]+", row):
                    continue
                cells = [value.strip() for value in row.strip("|").split("|")]
                if len(cells) != 3:
                    raise ValueError("The comparison table must contain exactly three columns")
                rows.append(cells)
            add_table(document, rows)
            for row_index, cells in enumerate(rows):
                if row_index == 0:
                    plain.append("[업무 처리 방식 비교]")
                else:
                    plain.extend([text_only(cells[0]), "현재(As-Is): " + text_only(cells[1]), "개선 후(To-Be): " + text_only(cells[2]), ""])
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(10)
            inline(paragraph, line[2:], size=17, all_bold=True)
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor.from_string("0C5551")
            plain.append(text_only(line[2:]))
        elif line.startswith("## "):
            title = line[3:]
            paragraph = document.add_paragraph(style="Heading 2")
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.keep_with_next = True
            if title.startswith("3."):
                paragraph.paragraph_format.page_break_before = True
            inline(paragraph, title, size=12, all_bold=True)
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor.from_string("172E3B")
            plain.append(title)
        elif line.startswith("- "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.3)
            paragraph.paragraph_format.first_line_indent = Cm(-0.3)
            paragraph.paragraph_format.space_after = Pt(3)
            inline(paragraph, "• " + line[2:])
            plain.append("• " + text_only(line[2:]))
        else:
            paragraph = document.add_paragraph()
            inline(paragraph, line)
            plain.append(text_only(line))
        i += 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = footer.add_run()
    set_font(run, size=8)
    run._r.extend((start, instruction, end))
    OUTPUT.mkdir(exist_ok=True)
    docx_path = OUTPUT / "contest-proposal.docx"
    txt_path = OUTPUT / "contest-proposal.txt"
    document.save(docx_path)
    txt_path.write_text("\n".join(plain).strip() + "\n", encoding="utf-8-sig")
    return docx_path, txt_path


if __name__ == "__main__":
    for artifact in build():
        print(artifact)
